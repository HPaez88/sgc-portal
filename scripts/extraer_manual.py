"""
extraer_manual.py — Convierte el Manual de la Calidad y sus anexos (formatos/Docus Sistema)
a backend/knowledge/manual/manual.json para consultarlos en el portal.

Uso:  python scripts/extraer_manual.py
Idempotente: re-genera el JSON y copia los originales a backend/knowledge/manual/archivos/.
"""
import json
import re
import shutil
import sys
from datetime import date
from pathlib import Path

import docx
import openpyxl
from docx.table import Table
from docx.text.paragraph import Paragraph

RAIZ = Path(__file__).resolve().parent.parent
ORIGEN = RAIZ / "formatos" / "Docus Sistema"
DESTINO = RAIZ / "backend" / "knowledge" / "manual"
ARCHIVOS = DESTINO / "archivos"

# clave, revisión y título a partir del nombre: "OOMLSC-A19  Rev 33  ANEXO 19. MATRIZ..."
RE_NOMBRE = re.compile(r"^(?P<clave>OOMLSC-[A-Z0-9]+)\s+Rev\s+(?P<rev>\d+)\s+(?P<titulo>.+)\.(?P<ext>\w+)$", re.I)
# encabezados del manual: "4.1 Comprensión...", "1.-Objetivo", "3-. Descripción General"
RE_SECCION = re.compile(r"^(?P<num>\d{1,2}(?:\.\d{1,2}){0,2})(?:\s*[.\-]+\s*|\s+)(?P<titulo>[A-ZÁÉÍÓÚÑ].{2,90})$")


def limpiar(texto: str) -> str:
    return re.sub(r"\s+", " ", texto or "").strip()


def meta_desde_nombre(nombre: str) -> dict:
    m = RE_NOMBRE.match(nombre)
    if not m:
        return {"clave": nombre, "revision": "", "titulo": nombre, "ext": nombre.rsplit(".", 1)[-1].lower()}
    titulo = re.sub(r"\.+$", "", limpiar(m["titulo"]))
    return {"clave": m["clave"].upper(), "revision": f"Rev. {m['rev']}", "titulo": titulo, "ext": m["ext"].lower()}


# ───────────────────────── Manual (docx) ─────────────────────────
def tabla_a_filas(tabla: Table) -> list[list[str]]:
    filas = []
    for fila in tabla.rows:
        celdas, previa = [], None
        for celda in fila.cells:
            t = limpiar(celda.text)
            if t == previa:  # celdas combinadas repiten el texto
                continue
            celdas.append(t)
            previa = t
        if any(celdas):
            filas.append(celdas)
    return filas


def extraer_manual(ruta: Path) -> dict:
    d = docx.Document(ruta)
    secciones = [{"numero": "0", "titulo": "Portada", "bloques": []}]
    for hijo in d.element.body.iterchildren():
        tag = hijo.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(hijo, d)
            texto = limpiar(p.text)
            if not texto:
                continue
            m = RE_SECCION.match(texto)
            es_titulo = bool(m) and not texto.endswith((":", ";")) and len(texto) < 100
            if es_titulo:
                secciones.append({"numero": m["num"], "titulo": limpiar(m["titulo"]), "bloques": []})
                continue
            tipo = "lista" if "List" in (p.style.name or "") else "p"
            secciones[-1]["bloques"].append({"tipo": tipo, "texto": texto})
        elif tag == "tbl":
            filas = tabla_a_filas(Table(hijo, d))
            if filas:
                secciones[-1]["bloques"].append({"tipo": "tabla", "filas": filas})
    secciones = [s for s in secciones if s["bloques"] or s["numero"] != "0"]  # los de nivel 1 pueden no tener texto propio
    return {"secciones": secciones, "total_secciones": len(secciones)}


# ───────────────────────── Anexos (xlsx) ─────────────────────────
def extraer_hoja(ws) -> dict | None:
    filas = []
    for fila in ws.iter_rows(values_only=True):
        celdas = ["" if v is None else limpiar(str(v)) for v in fila]
        if any(celdas):
            filas.append(celdas)
    if len(filas) < 3:
        return None
    # recortar columnas totalmente vacías
    ancho = max(len(f) for f in filas)
    usadas = [i for i in range(ancho) if any(i < len(f) and f[i] for f in filas)]
    filas = [[f[i] if i < len(f) else "" for i in usadas] for f in filas]
    # cabecera = primera fila con 3+ celdas llenas; lo anterior son metadatos
    idx = next((i for i, f in enumerate(filas) if sum(1 for c in f if c) >= 3), 0)
    meta = [" · ".join(c for c in f if c) for f in filas[:idx]]
    columnas = [c or f"Col {i + 1}" for i, c in enumerate(filas[idx])]
    datos = [f for f in filas[idx + 1:] if any(c and c != "#VALUE!" for c in f)]
    if not datos:
        return None
    return {"nombre": limpiar(ws.title), "meta": meta, "columnas": columnas, "filas": datos}


def extraer_xlsx(ruta: Path) -> dict:
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    hojas = [h for h in (extraer_hoja(ws) for ws in wb.worksheets) if h]
    return {"hojas": hojas}


# ───────────────────────── Orquestación ─────────────────────────
def main() -> int:
    if not ORIGEN.exists():
        print(f"No existe {ORIGEN}", file=sys.stderr)
        return 1
    ARCHIVOS.mkdir(parents=True, exist_ok=True)
    salida = {"generado": date.today().isoformat(), "manual": None, "anexos": []}
    for ruta in sorted(ORIGEN.iterdir()):
        if not ruta.is_file():
            continue
        meta = meta_desde_nombre(ruta.name)
        destino = ARCHIVOS / f"{meta['clave']}.{meta['ext']}"
        shutil.copyfile(ruta, destino)
        entrada = {**meta, "archivo": destino.name, "tamano_kb": round(ruta.stat().st_size / 1024)}
        try:
            if meta["ext"] == "docx" and meta["clave"] == "OOMLSC-01":
                salida["manual"] = {**entrada, **extraer_manual(ruta)}
                continue
            if meta["ext"] == "xlsx":
                entrada.update(tipo="tablas", **extraer_xlsx(ruta))
            elif meta["ext"] == "pdf":
                entrada.update(tipo="pdf")
            else:
                entrada.update(tipo="no_disponible",
                               nota=f"Formato .{meta['ext']} (Office 2003). Guardar como .{ 'xlsx' if meta['ext']=='xls' else 'docx'} para extraer el contenido.")
        except Exception as e:  # un anexo roto no debe tirar el resto
            entrada.update(tipo="error", nota=str(e))
        salida["anexos"].append(entrada)
    DESTINO.mkdir(parents=True, exist_ok=True)
    (DESTINO / "manual.json").write_text(json.dumps(salida, ensure_ascii=False, indent=1), encoding="utf-8")
    m = salida["manual"] or {}
    print(f"manual: {m.get('total_secciones', 0)} secciones | anexos: {len(salida['anexos'])} → {DESTINO / 'manual.json'}")
    for a in salida["anexos"]:
        print(f"  {a['clave']:12} {a['tipo']:14} {a.get('titulo','')[:50]}  hojas={len(a.get('hojas', []))}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
